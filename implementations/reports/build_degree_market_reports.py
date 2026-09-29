"""Build separate attributed ITM and BA market reports from the public-feed collector."""
from __future__ import annotations

import asyncio
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
import sys

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
BACKEND = ROOT / "syllabus-check" / "backend"
sys.path.insert(0, str(BACKEND))

from services.reports.market_alignment import JOB_SOURCES, PROGRAMS  # noqa: E402
from services.scraper.focused_market import collect_focused_market_jobs  # noqa: E402

OUTPUT_DIR = Path(__file__).resolve().parent
NAVY = "17365D"
PALE = "F4F7FA"
GRAY = "D9D9D9"
ORANGE = "C75B12"
WHITE = "FFFFFF"
BLACK = "000000"


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    node = props.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        props.append(node)
    node.set(qn("w:fill"), fill)


def border(cell):
    props = cell._tc.get_or_add_tcPr()
    borders = props.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        props.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), GRAY)
        borders.append(node)


def cell_text(cell, value, *, bold=False, color=BLACK, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(str(value))
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(color)


def table(doc, headers, rows, widths, numeric=()):
    result = doc.add_table(rows=1, cols=len(headers))
    result.alignment = WD_TABLE_ALIGNMENT.CENTER
    result.autofit = False
    for index, header in enumerate(headers):
        cell = result.rows[0].cells[index]
        cell.width = Inches(widths[index])
        shade(cell, NAVY)
        border(cell)
        cell_text(cell, header, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row_index, row in enumerate(rows):
        cells = result.add_row().cells
        for index, value in enumerate(row):
            cell = cells[index]
            cell.width = Inches(widths[index])
            shade(cell, WHITE if row_index % 2 == 0 else PALE)
            border(cell)
            cell_text(cell, value, align=WD_ALIGN_PARAGRAPH.CENTER if index in numeric else WD_ALIGN_PARAGRAPH.LEFT)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def configure(doc: Document, program_name: str):
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string("263238")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    title = doc.styles["Title"]
    title.font.name = "Aptos Display"
    title.font.size = Pt(27)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    for name, size in (("Heading 1", 17), ("Heading 2", 12.5)):
        style = doc.styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.keep_with_next = True
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run(f"SyllabusCheck  |  {program_name} Market Alignment Report")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string("6B7280")


def add_metrics(doc, metrics):
    result = doc.add_table(rows=1, cols=len(metrics))
    result.alignment = WD_TABLE_ALIGNMENT.CENTER
    result.autofit = False
    for index, (value, label) in enumerate(metrics):
        cell = result.rows[0].cells[index]
        cell.width = Inches(6.5 / len(metrics))
        shade(cell, PALE)
        border(cell)
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(str(value))
        r.bold = True
        r.font.size = Pt(21)
        r.font.color.rgb = RGBColor.from_string(ORANGE)
        p2 = cell.add_paragraph(label)
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.runs[0].font.size = Pt(8)


def build_report(program: str, all_jobs: list[dict]) -> Path:
    jobs = [job for job in all_jobs if job["program"] == program]
    program_name = PROGRAMS[program]["name"]
    families = Counter(job["role_family"] for job in jobs)
    titles = Counter(job["title"] for job in jobs)
    sources = Counter(job["source"] for job in jobs)
    skills: dict[str, set[str]] = defaultdict(set)
    categories = {}
    for job in jobs:
        for item in job["skills"]:
            skills[item["skill"]].add(f'{job["source"]}:{job["external_id"]}')
            categories[item["skill"]] = item["category"]

    doc = Document()
    configure(doc, program_name)
    doc.add_paragraph(f"{program_name} Market Alignment Report", style="Title")
    subtitle = doc.add_paragraph("Job-title demand, explicit skills, sources, and curriculum-readiness framework")
    subtitle.runs[0].bold = True
    subtitle.runs[0].font.size = Pt(13)
    subtitle.runs[0].font.color.rgb = RGBColor.from_string(NAVY)
    meta = doc.add_paragraph(f"Prepared for SyllabusCheck  |  {datetime.now(timezone.utc).strftime('%d %B %Y')}  |  Live public-feed sample")
    for run in meta.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string("6B7280")

    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph(
        f"This degree-specific report analyzes {len(jobs)} qualifying {program_name} job descriptions from "
        f"{len(sources)} active public feeds. The sample represents {len(titles)} distinct job titles across "
        f"{len(families)} role families. Skills are counted only when explicitly present in a job description."
    )
    doc.add_paragraph(
        "The evidence is suitable for identifying current role and skill signals in the collected sample. It does not "
        "predict hiring outcomes. Curriculum coverage remains not assessed until the applicable syllabi are uploaded and reviewed."
    )
    add_metrics(doc, [
        (len(jobs), "Qualified jobs"), (len(titles), "Distinct titles"),
        (len(families), "Role families"), (len(sources), "Active sources"),
    ])

    doc.add_heading("Job Families", level=1)
    table(doc, ["Role family", "Jobs", "Share"], [
        (name, count, f"{count * 100 / len(jobs):.1f}%") for name, count in families.most_common()
    ], [4.25, 1.0, 1.25], numeric=(1, 2))

    doc.add_page_break()
    doc.add_heading("Job Titles Represented", level=1)
    doc.add_paragraph(
        "The table preserves the employer-posted title rather than replacing it with a standardized label. "
        "This makes the report useful for student job searches while the role-family mapping supports aggregation."
    )
    table(doc, ["Job title", "Listings"], titles.most_common(30), [5.35, 1.15], numeric=(1,))

    doc.add_page_break()
    doc.add_heading("Skills Found in Job Descriptions", level=1)
    skill_rows = sorted(
        ((skill, categories[skill], len(job_ids), f"{len(job_ids) * 100 / len(jobs):.1f}%") for skill, job_ids in skills.items()),
        key=lambda row: (-row[2], row[0].lower()),
    )[:20]
    table(doc, ["Skill", "Category", "Jobs", "Demand share"], skill_rows, [2.2, 2.0, 0.9, 1.4], numeric=(2, 3))
    doc.add_heading("How to Read Skill Demand", level=2)
    doc.add_paragraph(
        "Jobs is the number of distinct descriptions that explicitly mention the skill. Demand share divides that count by all "
        "qualified jobs for this degree. A repeated term within one posting counts once, preventing long descriptions from receiving extra weight."
    )
    for text in (
        "High-frequency skills should be visible in learning outcomes and assessed work when they fit the academic purpose of the degree.",
        "Lower-frequency skills can still matter when they define a specialized role family or fast-growing technology area.",
        "Tool names should be interpreted with transferable concepts; for example, a cloud-platform requirement may represent architecture, security, deployment, and cost-management abilities.",
        "Communication and stakeholder skills require observable evidence such as presentations, requirements documents, decision logs, and team delivery artifacts.",
    ):
        doc.add_paragraph(text, style="List Bullet")

    doc.add_page_break()
    doc.add_heading("Job Sources and Attribution", level=1)
    source_rows = []
    for source, details in JOB_SOURCES.items():
        source_rows.append((details["name"], sources.get(source, 0), details["homepage"], details["api_url"]))
    table(doc, ["Source", "Jobs", "Website", "Public feed"], source_rows, [1.1, 0.55, 2.25, 2.6], numeric=(1,))
    doc.add_paragraph(
        "Each imported listing retains its canonical source URL. Source names and links must remain visible wherever listings or derived evidence are displayed."
    )
    doc.add_heading("Collection Architecture", level=2)
    table(doc, ["Stage", "Control"], [
        ("Acquire", "Request each public feed independently with timeouts and source-specific field mappings."),
        ("Normalize", "Convert HTML descriptions to text and standardize identifiers, dates, location, work type, and canonical URLs."),
        ("Qualify", "Require a substantive description and a title phrase mapped to a documented degree role family."),
        ("Extract", "Match maintained skill patterns against description text; never infer a skill solely from the title."),
        ("Deduplicate", "Use the source plus external identifier as the stable key so repeated collection updates rather than duplicates."),
        ("Attribute", "Retain the source name and canonical listing URL in the database, interface, and exported report."),
    ], [1.2, 5.3])
    doc.add_paragraph(
        "A zero in the source table means that the feed was configured but produced no degree-qualified listing in this collection. "
        "It should not be interpreted as proof that the source never carries relevant roles."
    )

    doc.add_page_break()
    doc.add_heading("Curriculum Alignment Status", level=1)
    doc.add_paragraph(
        "Status: not assessed. No authoritative degree syllabus evidence was available to score against these market skills. "
        "When syllabi are uploaded, each explicit skill will be classified as covered, partially covered, or missing and weighted by distinct job mentions."
    )
    doc.add_heading("Evidence Rubric", level=2)
    table(doc, ["Status", "Score", "Required evidence"], [
        ("Covered", "1.0", "An explicit learning outcome or assessed assignment requires the student to apply the skill."),
        ("Partial", "0.5", "The skill appears in topics or readings, but assessed application is limited or indirect."),
        ("Missing", "0.0", "No verified syllabus evidence addresses the skill."),
        ("Not assessed", "—", "The necessary syllabus or course mapping has not been uploaded and verified."),
    ], [1.2, 0.7, 4.6], numeric=(1,))
    doc.add_heading("Market-Weighted Calculation", level=2)
    doc.add_paragraph(
        "For every skill, multiply its coverage value by the number of distinct job descriptions mentioning it. Sum those values "
        "and divide by the total job-mention weight. This prioritizes repeatedly observed demand while retaining a traceable result for each skill."
    )
    doc.add_heading("Syllabus Evidence Required", level=2)
    for text in (
        "Course code, title, catalog description, and degree-program mapping.",
        "Learning outcomes, weekly topics, required tools, readings, and laboratory activities.",
        "Assessments and projects showing what students must produce or demonstrate.",
        "Faculty review of extracted topics before they are treated as authoritative coverage evidence.",
    ):
        doc.add_paragraph(text, style="List Bullet")

    doc.add_page_break()
    doc.add_heading("Student Job-Readiness Framework", level=1)
    doc.add_paragraph(
        "Curriculum alignment describes exposure and assessed learning. Student readiness requires additional evidence that the individual "
        "can apply those skills to relevant problems and communicate the result to an employer."
    )
    table(doc, ["Evidence layer", "Student evidence", "System use"], [
        ("Curriculum", "Completed courses, assessed work, and verified learning outcomes", "Shows where demanded skills are taught and practiced."),
        ("Portfolio", "Projects, datasets, systems, dashboards, reports, and measurable outcomes", "Demonstrates applied ability beyond a keyword claim."),
        ("Resume", "Degree, experience, projects, certifications, and quantified accomplishments", "Connects each claimed skill to supporting evidence."),
        ("Role fit", "Target title, preferred industries, location, work authorization, and seniority", "Filters jobs that are plausible and relevant."),
        ("Application", "Tailored resume, cover letter, interview examples, and networking activity", "Tracks preparation and execution for each opportunity."),
    ], [1.05, 2.7, 2.75])
    doc.add_heading("Degree-Specific Portfolio Priorities", level=2)
    priorities = (
        (
            "Deliver a scoped technology project with schedule, risks, stakeholder decisions, and measurable outcomes.",
            "Design or integrate a secure API or enterprise workflow and document the architecture.",
            "Deploy a cloud-based solution with identity, monitoring, cost, and governance considerations.",
            "Translate business requirements into a system design, test plan, and implementation recommendation.",
        ) if program == "itm" else (
            "Complete a reproducible SQL and Python analysis using a realistic business dataset.",
            "Build a decision-oriented dashboard and explain metric definitions, data quality, and limitations.",
            "Frame an ambiguous business question, document assumptions, and present an actionable recommendation.",
            "Show data preparation, statistical reasoning, validation, visualization, and stakeholder communication in one end-to-end project.",
        )
    )
    for text in priorities:
        doc.add_paragraph(text, style="List Bullet")

    doc.add_page_break()
    doc.add_heading("Implementation Roadmap", level=1)
    table(doc, ["Phase", "Timing", "Deliverables", "Exit condition"], [
        ("1. Stabilize evidence", "Weeks 1–2", "Restore database connection; run all feeds; verify source counts and canonical links.", "Repeatable import with recorded successes and failures."),
        ("2. Validate taxonomy", "Weeks 2–4", "Faculty and career-services review of degree roles, title phrases, and skill synonyms.", "Approved mapping with false-positive and false-negative samples."),
        ("3. Load curriculum", "Weeks 3–6", "Current syllabi, program mappings, course outcomes, topics, tools, and assessments.", "Every included course has an owner and review status."),
        ("4. Score alignment", "Weeks 5–8", "Coverage rows, weighted degree score, missing-skill priorities, and faculty review workflow.", "Scores are reproducible from stored evidence."),
        ("5. Add student evidence", "Weeks 7–10", "Private resume, project, course-completion, preference, and authorization records.", "Row-level access tests pass for each user role."),
        ("6. Pilot advising", "Weeks 10–12", "Advisor dashboard, student readiness view, role matches, and feedback collection.", "Pilot users can explain every recommendation and its evidence."),
    ], [1.15, 0.85, 2.8, 1.7])
    doc.add_heading("Operating Metrics", level=2)
    for text in (
        "Evidence volume: qualified jobs, active sources, distinct titles, and descriptions retained by degree and collection date.",
        "Quality: manually reviewed precision for role classification and explicit skill extraction.",
        "Coverage: percentage of priority skills with verified course evidence and assessed student work.",
        "Freshness: age of the newest and oldest active listing, plus source-specific collection failures.",
        "Student action: portfolio evidence completed, resume gaps resolved, suitable roles saved, and applications prepared.",
    ):
        doc.add_paragraph(text, style="List Bullet")

    doc.add_page_break()
    doc.add_heading("Method and Limitations", level=1)
    table(doc, ["Area", "Method or limitation"], [
        ("Qualification", "A substantive description and an explicit degree-relevant title phrase are required."),
        ("Skill demand", "A skill counts once per listing and only when present in description text."),
        ("Source resilience", "Public feeds are collected independently; one unavailable feed does not stop the remaining sources."),
        ("Sample scope", "Remote and European feeds are overrepresented; the sample is not a complete labor-market census."),
        ("Student outcome", "Alignment supports preparation planning and does not guarantee an interview, offer, salary, or employment."),
    ], [1.45, 5.05])
    doc.add_heading("Data Governance and Security", level=1)
    table(doc, ["Control", "Implementation expectation"], [
        ("Data minimization", "Store only the profile, academic, resume, and preference fields needed for matching and advising."),
        ("Private-by-default access", "Use Supabase authentication and row-level security so students see their own records and approved staff see only authorized scopes."),
        ("Service credentials", "Keep service-role keys on trusted servers; never expose them in the browser, reports, logs, or source control."),
        ("Auditability", "Record privileged reads and writes, source imports, report generation, role changes, and administrative actions."),
        ("Retention", "Define retention and deletion rules for resumes, parsed text, generated reports, and inactive accounts."),
        ("Human review", "Allow users and faculty to correct extracted evidence, mappings, and recommendations before consequential use."),
    ], [1.5, 5.0])
    doc.add_heading("Interpretation", level=1)
    doc.add_paragraph(
        "This report is a transparent snapshot of observed demand in the configured public feeds. It supports curriculum review, student preparation, "
        "and evidence-based advising. Decisions should combine this report with local employer input, longitudinal collections, faculty judgment, "
        "individual student evidence, and applicable institutional policy."
    )

    filename = "ITM_Market_Alignment_Report.docx" if program == "itm" else "Business_Analytics_Market_Alignment_Report.docx"
    output = OUTPUT_DIR / filename
    doc.save(output)
    return output


async def main():
    jobs = await collect_focused_market_jobs(pages=5)
    if not jobs:
        raise RuntimeError("No qualifying public-feed jobs were collected.")
    for program in ("itm", "ba"):
        print(build_report(program, jobs))


if __name__ == "__main__":
    asyncio.run(main())
